"""门店联系人导入与查询。联系人是精确业务数据，禁止写入 RAG。"""
from __future__ import annotations

import csv
import re
from io import BytesIO
from pathlib import Path
from typing import Any

from cryptography.fernet import Fernet
from sqlalchemy import select

from config.settings import settings
from database.models import Store, StoreContact

_HEADER_ALIASES = {
    "store_name": ("门店名称", "门店", "美团门店全名"),
    "branch": ("分店",),
    "manager_name": ("店长姓名", "店长"),
    "manager_mobile": ("店长电话", "店长手机号", "手机号", "电话"),
}
_PHONE_RE = re.compile(r"^1[3-9]\d{9}$")


def _normal(value: object) -> str:
    return re.sub(r"[\s·•,，。()（）\[\]【】_\-—]", "", str(value or "")).lower()


def _pick(row: dict[str, Any], aliases: tuple[str, ...]) -> str:
    for key in aliases:
        value = row.get(key)
        if value is not None and str(value).strip():
            return str(value).strip()
    return ""


def _mask_mobile(mobile: str) -> str:
    return f"{mobile[:3]}****{mobile[-4:]}" if len(mobile) >= 7 else "***"


def _read_rows(filename: str, content: bytes) -> list[dict[str, Any]]:
    suffix = Path(filename or "").suffix.lower()
    if suffix == ".csv":
        last_error: Exception | None = None
        for encoding in ("utf-8-sig", "utf-8", "gb18030", "gbk"):
            try:
                return list(csv.DictReader(content.decode(encoding).splitlines()))
            except UnicodeDecodeError as exc:
                last_error = exc
        raise ValueError(f"CSV 编码无法识别：{last_error}")
    if suffix in (".xlsx", ".xls"):
        if suffix == ".xls":
            raise ValueError("暂不支持 .xls，请另存为 .xlsx 或 .csv 后再导入")
        from openpyxl import load_workbook

        book = load_workbook(BytesIO(content), read_only=True, data_only=True)
        try:
            sheet = book.active
            values = list(sheet.iter_rows(values_only=True))
        finally:
            book.close()
        if not values:
            return []
        headers = [str(v or "").strip() for v in values[0]]
        return [
            {headers[i]: row[i] if i < len(row) else "" for i in range(len(headers))}
            for row in values[1:]
            if any(v not in (None, "") for v in row)
        ]
    raise ValueError("联系人文件仅支持 CSV 或 XLSX")


def _resolve_store(row: dict[str, Any], stores: list[Store]) -> Store | None:
    candidates = [
        *(str(row.get(key) or "").strip() for key in _HEADER_ALIASES["store_name"]),
        _pick(row, _HEADER_ALIASES["branch"]),
    ]
    scores: dict[int, tuple[int, Store]] = {}
    for raw in candidates:
        token = _normal(raw)
        if not token:
            continue
        for store in stores:
            name = _normal(store.store_name)
            keyword = _normal(store.search_keyword)
            # 完整美团门店名优先于简称。简称可同时命中多个品牌，不能因此错绑店长。
            score = 0
            if token == name:
                score = 100
            elif name and name in token:
                score = 90
            elif name and token in name:
                score = 70
            elif keyword and token == keyword:
                score = 60
            elif keyword and keyword in token:
                score = 50
            elif keyword and token in keyword:
                score = 40
            if score > scores.get(store.id, (0, store))[0]:
                scores[store.id] = (score, store)
    if not scores:
        return None
    best_score = max(score for score, _ in scores.values())
    best = [store for score, store in scores.values() if score == best_score]
    return best[0] if len(best) == 1 else None


def _validate_contact_import(filename: str, content: bytes, db_session) -> tuple[dict, list[dict]]:
    """返回可展示的预览和仅供当前进程导入用的原始行。"""
    rows = _read_rows(filename, content)
    stores = list(db_session.execute(select(Store)).scalars().all())
    preview: list[dict] = []
    valid: list[dict] = []
    seen_store_ids: set[int] = set()

    for index, row in enumerate(rows, start=2):
        manager_name = _pick(row, _HEADER_ALIASES["manager_name"])
        mobile = re.sub(r"\D", "", _pick(row, _HEADER_ALIASES["manager_mobile"]))
        store = _resolve_store(row, stores)
        errors: list[str] = []
        if store is None:
            errors.append("无法唯一匹配门店，请检查美团门店全名或分店字段")
        if not manager_name:
            errors.append("缺少店长姓名")
        if not _PHONE_RE.fullmatch(mobile):
            errors.append("店长电话必须是 11 位中国大陆手机号")
        if store and store.id in seen_store_ids:
            errors.append("同一门店在导入文件中重复")
        if store:
            seen_store_ids.add(store.id)
        item = {
            "row": index,
            "store_id": store.id if store else None,
            "store_name": store.store_name if store else _pick(row, _HEADER_ALIASES["store_name"]),
            "manager_name": manager_name,
            "manager_mobile": _mask_mobile(mobile) if mobile else "",
            "errors": errors,
        }
        preview.append(item)
        if not errors and store:
            valid.append({"store": store, "manager_name": manager_name, "mobile": mobile, "row": index})

    return {
        "success": True,
        "file": Path(filename or "").name,
        "total": len(rows),
        "valid": len(valid),
        "invalid": len(rows) - len(valid),
        "rows": preview,
    }, valid


def preview_contact_import(filename: str, content: bytes, db_session) -> dict:
    """解析并校验联系人文件；预览结果绝不返回完整手机号。"""
    preview, _ = _validate_contact_import(filename, content, db_session)
    return preview


def import_contacts(filename: str, content: bytes, db_session) -> dict:
    """按 store_id 幂等 upsert 有效联系人；无加密密钥时 fail-closed。"""
    if not settings.contact_encryption_key:
        raise ValueError("未配置 BIZ_CONTACT_ENCRYPTION_KEY，拒绝保存店长手机号")
    try:
        cipher = Fernet(settings.contact_encryption_key.encode("utf-8"))
    except Exception as exc:  # noqa: BLE001
        raise ValueError("BIZ_CONTACT_ENCRYPTION_KEY 不是有效 Fernet 密钥") from exc

    result, valid_rows = _validate_contact_import(filename, content, db_session)
    imported = 0
    for item in valid_rows:
        store = item["store"]
        contact = db_session.execute(
            select(StoreContact).where(StoreContact.store_id == store.id)
        ).scalars().first()
        encrypted = cipher.encrypt(item["mobile"].encode("utf-8")).decode("utf-8")
        values = {
            "store_name": store.store_name,
            "manager_name": item["manager_name"],
            "manager_mobile_encrypted": encrypted,
            "manager_mobile_last4": item["mobile"][-4:],
            "active": True,
            "source_file": Path(filename or "").name,
            "source_row": item["row"],
        }
        if contact:
            for key, value in values.items():
                setattr(contact, key, value)
        else:
            db_session.add(StoreContact(store_id=store.id, **values))
        imported += 1
    db_session.commit()
    result["imported"] = imported
    return result


def list_contacts(db_session) -> list[dict]:
    rows = db_session.execute(select(StoreContact).order_by(StoreContact.store_name)).scalars().all()
    return [
        {
            "store_id": r.store_id,
            "store_name": r.store_name,
            "manager_name": r.manager_name,
            "manager_mobile": f"***{r.manager_mobile_last4}",
            "dingtalk_user_id_configured": bool(r.dingtalk_user_id),
            "active": r.active,
            "source_file": r.source_file,
            "updated_at": str(r.updated_at)[:19],
        }
        for r in rows
    ]


def get_active_contact(store_id: int, db_session) -> StoreContact | None:
    return db_session.execute(
        select(StoreContact).where(StoreContact.store_id == store_id, StoreContact.active.is_(True))
    ).scalars().first()


def decrypt_contact_mobile(contact: StoreContact) -> str:
    """仅 live Dispatcher 可在内存中解密号码；不得返回给接口或日志。"""
    if not settings.contact_encryption_key:
        raise ValueError("未配置 BIZ_CONTACT_ENCRYPTION_KEY，拒绝解析钉钉收件人")
    try:
        cipher = Fernet(settings.contact_encryption_key.encode("utf-8"))
        return cipher.decrypt(contact.manager_mobile_encrypted.encode("utf-8")).decode("utf-8")
    except Exception as exc:  # noqa: BLE001
        raise ValueError("店长手机号解密失败，拒绝真实发送") from exc
