# -*- coding: utf-8 -*-
"""CSV/XLSX 知识文件解析测试。"""
from datetime import date

from openpyxl import Workbook

from rag.loader import ALLOWED_UPLOAD_EXTS, load_documents, split_documents_hierarchical


def test_csv_loader_supports_chinese_encoding_and_header_mapping(tmp_path):
    path = tmp_path / "店长信息库.csv"
    path.write_bytes(
        "门店名称,店长,联系电话\n一号店,张三,13800000000\n二号店,李四,13900000000\n".encode("gb18030")
    )

    docs = load_documents([path])

    assert len(docs) == 2
    assert docs[0].metadata["table_format"] == "csv"
    assert docs[0].metadata["sheet_name"] == "店长信息库"
    assert docs[0].metadata["table_row"] == 2
    assert "字段：门店名称 | 店长 | 联系电话" in docs[0].page_content
    assert "第2行：门店名称=一号店 | 店长=张三 | 联系电话=13800000000" in docs[0].page_content
    assert "第3行：门店名称=二号店 | 店长=李四 | 联系电话=13900000000" in docs[1].page_content


def test_xlsx_loader_reads_each_non_empty_sheet(tmp_path):
    path = tmp_path / "主题话术库.xlsx"
    workbook = Workbook()
    first = workbook.active
    first.title = "售后话术"
    first.append(["主题", "话术"])
    first.append(["退款", "您好，我来帮您处理退款。"])
    second = workbook.create_sheet("空白页")
    second.append([None, None])
    third = workbook.create_sheet("节日话术")
    third.append(["节日", "开场话术", "日期"])
    third.append(["春节", "新年快乐，欢迎光临。", date(2026, 2, 17)])
    workbook.save(path)
    workbook.close()

    docs = load_documents([path])
    chunks = split_documents_hierarchical(docs)

    assert [doc.metadata["sheet_name"] for doc in docs] == ["售后话术", "节日话术"]
    assert all(doc.metadata["table_format"] == "xlsx" for doc in docs)
    assert len({chunk.metadata["parent_id"] for chunk in chunks}) == len(docs)
    assert "主题=退款 | 话术=您好，我来帮您处理退款。" in docs[0].page_content
    assert "节日=春节 | 开场话术=新年快乐，欢迎光临。 | 日期=2026-02-17" in docs[1].page_content


def test_upload_policy_includes_spreadsheet_formats():
    assert ".csv" in ALLOWED_UPLOAD_EXTS
    assert ".xlsx" in ALLOWED_UPLOAD_EXTS
